"""手寫筆記產生器：把解答內容渲染成橫線筆記紙上的手寫頁面（PNG），再組成 Word 檔。

用法：
    python hwnotes.py content.py 輸出.docx [--png-dir 資料夾] [--seed N]

content.py 只要定義 NOTES（見 SKILL.md）。本檔的所有函式（mirror_diagram、
lens_diagram、wobbly_line、label…）都會自動放進 content.py 的全域，不用 import。
"""
import argparse, math, os, random, re, runpy, sys, tempfile
from PIL import Image, ImageDraw, ImageFont

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FONT = os.path.join(SKILL_DIR, 'fonts', 'ChenYuluoyan-nohint.ttf')

DPI = 200
PW, PH = round(19.0 / 2.54 * DPI), round(27.2 / 2.54 * DPI)   # A4 扣 1 cm 邊界
S = 44                # 基本字高 px
RULE = 76             # 橫線間距
TOP = 200             # 第一條橫線 y
MARGIN_X = 120        # 紅色邊線
TEXT_X = 150
RIGHT = PW - 70
FIG_X = TEXT_X + 120  # 圖的左上角 x；圖可用寬度約 FIG_W
FIG_W = RIGHT - FIG_X
FIG_ROWS = 5          # 光路圖預設佔 5 行
BLUE = (28, 48, 115)
RED = (190, 40, 40)
GRAY = (90, 90, 90)
PAPER = (253, 252, 246)
INDENT = [0, 70, 120]
NUM_W = 90            # 題號圈佔的寬度
ANS_PAD = 150         # 「Ans:」與框線佔的寬度

# 字型裡沒有的符號，自動換成手寫時常見的寫法
SUBS = {'≈': '=', '−': '-', '≥': '>=', '≤': '<=', '✓': 'OK', '⇒': '→', '∝': '~', '–': '-', '—': '-'}

_fonts, _cmap, _missing = {}, None, set()

def font(size):
    size = max(8, int(size))
    if size not in _fonts:
        _fonts[size] = ImageFont.truetype(FONT, size)
    return _fonts[size]

def _has_glyph(ch):
    global _cmap
    if _cmap is None:
        from fontTools.ttLib import TTFont
        _cmap = TTFont(FONT).getBestCmap()
    return ord(ch) in _cmap

# ---- 標記解析：^{上標} _{下標} [r]紅字[/r] ----
def parse(text, color=BLUE):
    for k, v in SUBS.items():
        text = text.replace(k, v)
    runs, mode, col, i, stack = [], 'n', color, 0, []
    while i < len(text):
        if text.startswith('[r]', i):
            col = RED; i += 3; continue
        if text.startswith('[/r]', i):
            col = color; i += 4; continue
        if text[i] in '^_' and i + 1 < len(text) and text[i + 1] == '{':
            stack.append(mode); mode = 'sup' if text[i] == '^' else 'sub'; i += 2; continue
        if text[i] == '}' and stack:
            mode = stack.pop(); i += 1; continue
        if text[i] != ' ' and not _has_glyph(text[i]):
            _missing.add(text[i])
        runs.append((text[i], mode, col)); i += 1
    return runs

def glyph_metrics(ch, mode, scale=1.0):
    size = S * scale * (0.6 if mode != 'n' else 1.0)
    adv = font(size).getlength(ch)
    if ord(ch) > 0x2E80:
        adv *= 0.95
    elif ch == ' ':
        adv *= 0.8
    return size, adv

def measure(runs, scale=1.0):
    return sum(glyph_metrics(c, m, scale)[1] for c, m, _ in runs)

def wrap(runs, width, scale=1.0):
    lines, cur, w = [], [], 0
    for r in runs:
        adv = glyph_metrics(r[0], r[1], scale)[1]
        if w + adv > width and cur:
            # 盡量在空白處斷行
            cut = max((k for k, x in enumerate(cur) if x[0] == ' '), default=-1)
            if cut > len(cur) * 0.5:
                lines.append(cur[:cut]); cur = cur[cut + 1:]
            else:
                lines.append(cur); cur = []
            w = measure(cur, scale)
        cur.append(r); w += adv
    if cur:
        lines.append(cur)
    return lines

def draw_runs(img, runs, x, base, scale=1.0, slope=None):
    """逐字貼上，每個字有隨機旋轉/大小/位移，基線緩慢飄移。回傳 (x_end, 範圍框)"""
    slope = random.uniform(-0.006, 0.006) if slope is None else slope
    phase = random.uniform(0, 6.28)
    x0 = x; top = base; bot = base
    for ch, mode, col in runs:
        size, adv = glyph_metrics(ch, mode, scale)
        adv *= random.uniform(0.93, 1.03)
        if ch != ' ':
            sz = size * random.uniform(0.95, 1.05)
            f = font(sz)
            pad = int(sz * 0.6)
            tile = Image.new('L', (int(sz * 2.2) + pad, int(sz * 2.2) + pad), 0)
            d = ImageDraw.Draw(tile)
            asc = f.getmetrics()[0]
            d.text((pad, pad), ch, font=f, fill=random.randint(215, 255),
                   stroke_width=1 if sz > 30 else 0, stroke_fill=random.randint(150, 200))
            tile = tile.rotate(random.uniform(-3.5, 3.5), resample=Image.BICUBIC, center=(pad + sz / 2, pad + asc * 0.7))
            dy = {'n': 0, 'sup': -S * scale * 0.42, 'sub': S * scale * 0.18}[mode]
            drift = 3 * math.sin(phase + (x - x0) / 160) + slope * (x - x0)
            px = int(x - pad + random.uniform(-1, 1))
            py = int(base - asc - pad + dy + drift + random.uniform(-2, 2))
            ink = Image.new('RGB', tile.size, tuple(min(255, max(0, c + random.randint(-12, 12))) for c in col))
            img.paste(ink, (px, py), tile)
            top = min(top, py + pad); bot = max(bot, py + pad + sz)
        x += adv
    return x, (x0, top, x, bot)

# ---- 手繪線條工具（給自訂圖用） ----
def label(img, text, x, y, col=BLUE, scale=0.8):
    """以 (x, y) 為中心寫一小段字，支援 ^{} _{} 標記"""
    runs = parse(text, col)
    w = measure(runs, scale)
    draw_runs(img, runs, x - w / 2, y + S * scale * 0.35, scale, slope=0)

def wobbly_line(d, p, q, col=BLUE, w=3):
    n = max(2, int(math.dist(p, q) / 25))
    pts = []
    for k in range(n + 1):
        t = k / n
        j = 0 if k in (0, n) else random.uniform(-1.6, 1.6)
        pts.append((p[0] + (q[0] - p[0]) * t + j * 0.3, p[1] + (q[1] - p[1]) * t + j))
    d.line(pts, fill=col, width=w, joint='curve')

def dashed_line(d, p, q, col=BLUE, w=2, dash=14):
    L = math.dist(p, q)
    n = max(1, int(L / dash))
    for k in range(0, n, 2):
        a, b = k / n, min(1, (k + 1) / n)
        wobbly_line(d, (p[0] + (q[0] - p[0]) * a, p[1] + (q[1] - p[1]) * a),
                    (p[0] + (q[0] - p[0]) * b, p[1] + (q[1] - p[1]) * b), col, w)

def arrow_head(d, tip, ang, col=BLUE, size=18):
    """ang 為箭頭指向（弧度，影像座標，-pi/2 = 向上）"""
    for s in (2.6, -2.6):
        a = ang + s
        d.line([tip, (tip[0] + size * math.cos(a), tip[1] + size * math.sin(a))], fill=col, width=3)

def ray(d, p, q, col=BLUE, w=3):
    """畫光線並在中段加箭頭"""
    wobbly_line(d, p, q, col, w)
    arrow_head(d, ((p[0] + q[0]) / 2, (p[1] + q[1]) / 2), math.atan2(q[1] - p[1], q[0] - p[0]), col)

def arc(d, c, r, a0, a1, col=BLUE):
    """角度弧（度，影像座標）"""
    pts = [(c[0] + r * math.cos(math.radians(a0 + (a1 - a0) * k / 12)),
            c[1] + r * math.sin(math.radians(a0 + (a1 - a0) * k / 12))) for k in range(13)]
    d.line(pts, fill=col, width=2, joint='curve')

def vec_arrow(d, p, q, col=BLUE, w=3):
    """一般箭頭（力、場、向量）"""
    wobbly_line(d, p, q, col, w)
    arrow_head(d, q, math.atan2(q[1] - p[1], q[0] - p[0]), col)

def dot(d, p, col=BLUE, r=5):
    d.ellipse([p[0] - r, p[1] - r, p[0] + r, p[1] + r], fill=col)

def wobbly_box(d, box, col):
    x0, y0, x1, y1 = box
    o = lambda: random.uniform(-4, 4)
    wobbly_line(d, (x0 + o() - 6, y0 + o()), (x1 + o() + 8, y0 + o()), col)
    wobbly_line(d, (x1 + o(), y0 + o() - 6), (x1 + o(), y1 + o() + 6), col)
    wobbly_line(d, (x1 + o() + 6, y1 + o()), (x0 + o() - 8, y1 + o()), col)
    wobbly_line(d, (x0 + o(), y1 + o() + 6), (x0 + o(), y0 + o() - 8), col)

def wobbly_circle(d, cx, cy, r, col):
    pts, start = [], random.uniform(0, 6.28)
    for k in range(48):
        a = start + k / 44 * 2 * math.pi            # 多畫一點，頭尾交疊像手畫
        rr = r * (1 + 0.06 * math.sin(a * 2 + start)) + random.uniform(-1, 1)
        pts.append((cx + rr * math.cos(a), cy + rr * 0.9 * math.sin(a)))
    d.line(pts, fill=col, width=3, joint='curve')

def _extend(P, Q, box):
    """從 P 沿 P→Q 方向延伸到框 (xmin, xmax, ymin, ymax) 邊界"""
    xmin, xmax, ymin, ymax = box
    dx, dy = Q[0] - P[0], Q[1] - P[1]
    ts = []
    if dx < 0: ts.append((xmin - P[0]) / dx)
    if dx > 0: ts.append((xmax - P[0]) / dx)
    if dy < 0: ts.append((ymin - P[1]) / dy)
    if dy > 0: ts.append((ymax - P[1]) / dy)
    t = min(x for x in ts if x > 0)
    return (P[0] + dx * t, P[1] + dy * t)

def _object_arrow(img, d, ox, ay, h):
    wobbly_line(d, (ox, ay), (ox, ay - h), BLUE, 5)
    arrow_head(d, (ox, ay - h), -math.pi / 2, BLUE, 22)
    label(img, '物', ox - 30, ay - h / 2)

def _image_arrow(img, d, ix, ay, hi, virtual):
    up = -math.pi / 2 if hi > 0 else math.pi / 2
    if virtual:
        dashed_line(d, (ix, ay), (ix, ay - hi), RED, 4, 10)
        arrow_head(d, (ix, ay - hi), up, RED, 18)
        if abs(hi) > 80:                             # 像夠高：標在箭頭旁，免得超出圖框壓到上一行
            label(img, '像(虛)', ix + 70, ay - hi / 2, RED)
        else:
            label(img, '像(虛)', ix + 10, ay - hi - 35 if hi > 0 else ay - hi + 35, RED)
    else:
        wobbly_line(d, (ix, ay), (ix, ay - hi), RED, 4)
        arrow_head(d, (ix, ay - hi), up, RED, 18)
        label(img, '像', ix + 40, ay - hi + 30 if hi < 0 else ay - hi - 30, RED)

# ---- 光路圖（像素單位；幾何照比例，所以光線交點就是公式算出的像） ----
def mirror_diagram(f, p, h, Vx_off=560, rows=FIG_ROWS):
    """面鏡：f>0 凹面鏡、f<0 凸面鏡，鏡面朝左；物體在鏡左 p 處、高 h。
    畫兩條主光線（平行主軸、射向頂點），近軸近似。p == f 時反射光平行不成像。"""
    def draw(img, d, x0, y0):
        Vx, ay = x0 + Vx_off, y0 + 190
        box = (x0 - 40, x0 + Vx_off + 320, y0 + 15, y0 + rows * RULE - 15)
        wobbly_line(d, (box[0], ay), (box[1], ay), GRAY, 2)          # 主軸
        Rv, H = 520, 165
        sgn = -1 if f > 0 else 1                                     # 凹面鏡兩端彎向左
        pts = [(Vx + sgn * (Rv - math.sqrt(Rv * Rv - y * y)), ay + y) for y in range(-H, H + 1, 10)]
        d.line(pts, fill=BLUE, width=4, joint='curve')
        for (mx, my) in pts[::3]:                                    # 鏡背斜線
            d.line([(mx + 3, my), (mx + 20, my - 15)], fill=BLUE, width=2)
        for name, xx in (('F', Vx - f), ('C', Vx - 2 * f)):
            if box[0] < xx < box[1]:
                dot(d, (xx, ay)); label(img, name, xx, ay + 32)
        ox = Vx - p
        _object_arrow(img, d, ox, ay, h)
        T = (ox, ay - h)
        P1 = (Vx + sgn * (Rv - math.sqrt(Rv * Rv - h * h)), ay - h)  # 打在畫出來的鏡面上
        ray(d, T, P1)
        F = (Vx - f, ay)
        Q1 = F if f > 0 else (2 * P1[0] - F[0], 2 * P1[1] - F[1])
        ray(d, P1, _extend(P1, Q1, box))
        V = (Vx, ay)
        ray(d, T, V)
        ray(d, V, _extend(V, (ox, ay + h), box))
        if abs(p - f) < 1:
            return
        q = 1 / (1 / f - 1 / p)
        hi = -q / p * h
        ix = Vx - q
        if q < 0:                                    # 虛像：反射光往鏡後延長
            dashed_line(d, P1, (ix, ay - hi)); dashed_line(d, V, (ix, ay - hi))
        _image_arrow(img, d, ix, ay, hi, q < 0)
    return draw

def lens_diagram(f, p, h, Lx_off=520, rows=FIG_ROWS):
    """薄透鏡：f>0 凸透鏡、f<0 凹透鏡；物體在鏡左 p 處、高 h。
    畫兩條主光線（平行主軸、通過光心）。p == f 時折射光平行不成像。"""
    def draw(img, d, x0, y0):
        Lx, ay = x0 + Lx_off, y0 + 190
        box = (x0 - 40, x0 + FIG_W - 20, y0 + 15, y0 + rows * RULE - 15)
        wobbly_line(d, (box[0], ay), (box[1], ay), GRAY, 2)
        H = 165
        wobbly_line(d, (Lx, ay - H), (Lx, ay + H), BLUE, 4)          # 透鏡用直線加箭頭表示
        for y, s in ((ay - H, -1), (ay + H, 1)):
            arrow_head(d, (Lx, y), s * math.pi / 2 if f > 0 else -s * math.pi / 2, BLUE, 20)
        af = abs(f)
        for name, xx in (('F', Lx - af), ('F′', Lx + af), ('2F', Lx - 2 * af), ('2F′', Lx + 2 * af)):
            if box[0] < xx < box[1]:
                dot(d, (xx, ay)); label(img, name, xx, ay + 32)
        ox = Lx - p
        _object_arrow(img, d, ox, ay, h)
        T = (ox, ay - h)
        P1 = (Lx, ay - h)
        ray(d, T, P1)
        Fp = (Lx + f, ay)                                            # 凸透鏡在右、凹透鏡在左
        Q1 = Fp if f > 0 else (2 * P1[0] - Fp[0], 2 * P1[1] - Fp[1])
        ray(d, P1, _extend(P1, Q1, box))
        O = (Lx, ay)
        ray(d, T, O)
        ray(d, O, _extend(O, (2 * Lx - ox, ay + h), box))
        if abs(p - f) < 1:
            return
        q = 1 / (1 / f - 1 / p)
        hi = -q / p * h
        ix = Lx + q
        if q < 0:                                    # 虛像：折射光往左延長
            dashed_line(d, P1, (ix, ay - hi)); dashed_line(d, O, (ix, ay - hi))
        _image_arrow(img, d, ix, ay, hi, q < 0)
    return draw

# ---- 排版 ----
def new_page():
    img = Image.new('RGB', (PW, PH), PAPER)
    d = ImageDraw.Draw(img)
    y = TOP
    while y < PH - 40:
        d.line([(0, y), (PW, y)], fill=(196, 214, 236), width=2)
        y += RULE
    d.line([(MARGIN_X, 0), (MARGIN_X, PH)], fill=(232, 150, 150), width=2)
    return img

def line_width(it):
    return RIGHT - TEXT_X - INDENT[it[2]] - (ANS_PAD if it[0] == 'a' else 0)

def rows_of(it):
    k = it[0]
    if k == 'gap': return 1
    if k in ('h', 'brk'): return 0          # 標題寫在第一條線上方
    if k == 'fig': return it[2]
    if k == 'q': return len(wrap(parse(it[2]), RIGHT - TEXT_X - NUM_W))
    return len(wrap(parse(it[1]), line_width(it)))

SUBPART = re.compile(r'^\s*\(([b-z])\)')

def layout(notes):
    """排成頁：一題盡量不跨頁；一題比一頁還長時在 (b)(c)… 小題前斷開；('brk',) 強制換頁"""
    per_page = (PH - TOP - 40) // RULE
    blocks, cur = [], []
    for it in notes:
        if it[0] in ('q', 'h', 'brk') and cur:
            blocks.append(cur); cur = []
        cur.append(it)
    blocks.append(cur)
    pieces = []
    for b in blocks:
        if sum(rows_of(x) for x in b) <= per_page:
            pieces.append(b); continue
        part = []
        for it in b:
            if it[0] == 't' and SUBPART.match(it[1]) and part:
                pieces.append(part); part = []
            part.append(it)
        pieces.append(part)
    pages, page, used = [], [], 0
    for b in pieces:
        n = sum(rows_of(it) for it in b)
        if page and (b[0][0] == 'brk' or used + n > per_page):
            pages.append(page); page, used = [], 0
            while b and b[0][0] in ('gap', 'brk'):
                n -= rows_of(b[0]); b = b[1:]
        page.extend(b); used += n
        if n > per_page:
            print(f'警告：有一段 {n} 行，超過一頁 {per_page} 行，底部會被截掉', file=sys.stderr)
    pages.append(page)
    return pages

def render_page(items):
    img = new_page()
    d = ImageDraw.Draw(img)
    row = 0
    def base(r): return TOP + r * RULE - 10
    for it in items:
        kind = it[0]
        if kind == 'brk':
            continue
        if kind == 'gap':
            row += 1; continue
        if kind == 'h':
            x_end, _ = draw_runs(img, parse(it[1]), TEXT_X, TOP - 30, scale=1.25)
            wobbly_line(d, (TEXT_X - 5, TOP - 14), (x_end + 10, TOP - 18), BLUE, 3)
            continue
        if kind == 'fig':
            it[1](img, d, FIG_X, TOP + row * RULE)
            row += it[2]; continue
        if kind == 'q':
            num, text = it[1], it[2]
            b = base(row + 1)
            nr = parse(num, RED)
            w = measure(nr, 1.1)
            draw_runs(img, nr, TEXT_X + 30 - w / 2, b, scale=1.1, slope=0)
            wobbly_circle(d, TEXT_X + 30, b - S * 0.38, max(30, w / 2 + 14), RED)
            for line in wrap(parse(text), RIGHT - TEXT_X - NUM_W):
                row += 1
                draw_runs(img, line, TEXT_X + NUM_W, base(row))
            continue
        lines = wrap(parse(it[1], RED if kind == 'a' else BLUE), line_width(it))
        for k, line in enumerate(lines):
            row += 1
            x = TEXT_X + INDENT[it[2]]
            if kind == 'a' and k == 0:
                x, _ = draw_runs(img, parse('Ans: ', RED), x, base(row))
            x_end, box = draw_runs(img, line, x + (6 if kind == 'a' else 0), base(row))
            if kind == 'a':
                wobbly_box(d, (box[0] - 8, base(row) - S * 0.95, box[2] + 6, base(row) + 14), RED)
    return img

def build(notes, out_docx, png_dir):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    pngs = []
    for i, items in enumerate(layout(notes), 1):
        p = os.path.join(png_dir, f'page{i}.png')
        render_page(items).save(p, dpi=(DPI, DPI))
        pngs.append(p)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(1.0)
    st = doc.styles['Normal']
    st.font.size = Pt(12)
    st.paragraph_format.space_before = st.paragraph_format.space_after = Pt(0)
    for i, p in enumerate(pngs):
        para = doc.paragraphs[0] if i == 0 and doc.paragraphs else doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        run.add_picture(p, width=Cm(19.0))
        if i < len(pngs) - 1:
            run.add_break(WD_BREAK.PAGE)
    doc.save(out_docx)
    return pngs

def main():
    ap = argparse.ArgumentParser(description='手寫筆記 Word 產生器')
    ap.add_argument('content', help='定義 NOTES 的 .py 檔')
    ap.add_argument('out', help='輸出 .docx 路徑')
    ap.add_argument('--png-dir', help='保留每頁 PNG 的資料夾（檢查用）；不給就用暫存並刪掉')
    ap.add_argument('--seed', type=int, default=1, help='筆跡隨機種子，換數字筆跡就不同')
    a = ap.parse_args()
    if not os.path.exists(FONT):
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        from get_font import ensure_font
        ensure_font()
    random.seed(a.seed)
    ns = runpy.run_path(a.content, init_globals={k: v for k, v in globals().items() if not k.startswith('__')})
    notes = ns['NOTES']
    if a.png_dir:
        os.makedirs(a.png_dir, exist_ok=True)
        pngs = build(notes, a.out, a.png_dir)
    else:
        with tempfile.TemporaryDirectory() as tmp:
            pngs = build(notes, a.out, tmp)
    if _missing:
        print('警告：字型缺字（會畫成空白），請改寫法：' + ' '.join(sorted(_missing)), file=sys.stderr)
    print(f'{a.out}  {len(pngs)} 頁')
    if a.png_dir:
        for p in pngs:
            print(p)

if __name__ == '__main__':
    main()
